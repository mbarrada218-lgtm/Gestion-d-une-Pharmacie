from openpyxl import load_workbook, Workbook
from datetime import datetime
import os

FICHIER_EXCEL = "MEDICAMENTS.xlsx"
SEUIL_STOCK_FAIBLE = 10

def medicament_controlleur(tableau_general):
    if os.path.exists(FICHIER_EXCEL):
        wb = load_workbook(FICHIER_EXCEL)
        ws = wb["Medicaments"]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Medicaments"

        ws["A1"] = "N°"
        ws["B1"] = "Nom du medicament"
        ws["C1"] = "Dossage (ml/mg)"
        ws["D1"] = "Categorie"
        ws["E1"] = "Quantite en stock"
        ws["F1"] = "Date expiration"
        ws["G1"] = "Prix d'achat"
        ws["H1"] = "Prix de vente"
        ws["I1"] = "Fournisseur"

        wb.save(FICHIER_EXCEL)

    for row in ws.iter_rows(min_row=2, values_only=True):
        tableau_general.insert("", "end", values=row)

    return wb, ws


def validation_des_donnees(nom, dosage, quantite, date_ex, categorie, p_achat, p_vente, fourni):

    if not nom or not dosage or not quantite or not date_ex or not categorie or not p_achat or not p_vente or not fourni:
        return False, "Remplis tous les champs !"

    try:
        dosage = float(dosage)
        quantite = int(quantite)
        p_achat = float(p_achat)
        p_vente = float(p_vente)
    except ValueError:
        return False, "Le dosage, la quantité et les prix doivent être des nombres !"

    try:
        datetime.strptime(date_ex, "%d/%m/%Y")
    except ValueError:
        return False, "La date d'expiration doit être au format JJ/MM/AAAA !"

    valeurs = {
        "nom": nom,
        "dosage": dosage,
        "quantite": quantite,
        "date_ex": date_ex,
        "categorie": categorie,
        "p_achat": p_achat,
        "p_vente": p_vente,
        "fourni": fourni,
    }
    return True, valeurs


def enregistrer_medicament(wb, ws, tableau_general, nom, dosage, quantite,
                            date_ex, categorie, p_achat, p_vente, fourni):
    """
    Valide les champs saisis dans la vue, puis enregistre le medicament
    dans le fichier Excel et dans tableau_general.
    Retourne (True, None) si l'enregistrement a reussi,
    sinon (False, message_erreur) pour que la vue affiche un messagebox.
    """
    ok, resultat = validation_des_donnees(nom, dosage, quantite, date_ex,
                                           categorie, p_achat, p_vente, fourni)
    if not ok:
        return False, resultat

    v = resultat
    numero = len(tableau_general.get_children()) + 1
    valeurs = (numero, v["nom"], v["dosage"], v["categorie"], v["quantite"],
               v["date_ex"], v["p_achat"], v["p_vente"], v["fourni"])

    # 1) Ajouter dans Excel
    ws.append(valeurs)
    wb.save(FICHIER_EXCEL)

    # 2) Ajouter dans le tableau
    tableau_general.insert("", "end", values=valeurs)

    return True, None


# ============================================================================
#                 ALERTES : MEDICAMENTS EXPIRES / STOCK FAIBLE
# ============================================================================

def charger_medicaments_expires(ws, tableau_medicament):

    for item in tableau_medicament.get_children():
        tableau_medicament.delete(item)

    aujourdhui = datetime.now()

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or row[5] is None:
            continue

        date_ex = row[5]
        try:
            if isinstance(date_ex, str):
                date_expiration = datetime.strptime(date_ex, "%d/%m/%Y")
            else:
                date_expiration = date_ex  # deja un objet date/datetime (Excel)
        except (ValueError, TypeError):
            continue

        if date_expiration < aujourdhui:
            tableau_medicament.insert("", "end", values=row)


def charger_stock_faible(ws, tableau_stock, seuil=SEUIL_STOCK_FAIBLE):

    for item in tableau_stock.get_children():
        tableau_stock.delete(item)

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or row[4] is None:
            continue
        try:
            quantite = int(row[4])
        except (ValueError, TypeError):
            continue
        if quantite <= seuil:
            tableau_stock.insert("", "end", values=row)