from openpyxl import load_workbook, Workbook
from datetime import datetime
from tkinter import messagebox
import os

FICHIER_EXCEL = "VENTS.xlsx"


def vents_controlleur(tableau_historique):

    if os.path.exists(FICHIER_EXCEL):
        wb = load_workbook(FICHIER_EXCEL)
        ws = wb["VENTS"]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "VENTS"

        ws["A1"] = "ID"
        ws["B1"] = "N° SERIE"
        ws["C1"] = "Date"
        ws["D1"] = "Nom du Médicament"
        ws["E1"] = "Dosage (mg/ml)"
        ws["F1"] = "Prix"
        ws["G1"] = "Quantité"
        ws["H1"] = "Médecin"
        ws["I1"] = "N° Ordonnance"
        ws["J1"] = "Nom du Client"

        wb.save(FICHIER_EXCEL)

    # Charger les anciennes ventes dans Treeview

    for row in ws.iter_rows(min_row=2, values_only=True):
        tableau_historique.insert("","end",values=row)
    return wb, ws


def validation_donnees(RECH,NOM,DOSAGE,STOCK,PR_VENT,QUAN,NOM_CLIENT,MEDECIN,
                       N_ORD,wb,ws,tableau_historique):

    nom = NOM.get()
    dosage = DOSAGE.get()
    stock = STOCK.get()
    pr_vent = PR_VENT.get()
    quant = QUAN.get()
    nom_client = NOM_CLIENT.get()
    medecin = MEDECIN.get()
    ordonnance = N_ORD.get()

    # Vérification بسيطة
    if not nom or not quant:
        messagebox.showwarning("Attention","Remplis les informations nécessaires")
        return
    # Nouvelle ligne
    def generer_numero_serie(ws):

        dernier_numero = ws.max_row - 1

        nouveau_numero = dernier_numero + 1

        return f"MED{nouveau_numero:06d}"
    
    numero_serie = generer_numero_serie(ws)
    nouvelle_ligne = ws.max_row + 1

    # ID
    ws["A" + str(nouvelle_ligne)] = nouvelle_ligne - 1
    # N° SERIE
    # إلا عندك Entry ديال numéro série خاصك تدوزو للفونكسيون
    ws["B" + str(nouvelle_ligne)] = numero_serie
    # Date
    ws["C" + str(nouvelle_ligne)] = datetime.now().strftime("%d/%m/%Y")
    # Nom médicament
    ws["D" + str(nouvelle_ligne)] = nom
    # Dosage
    ws["E" + str(nouvelle_ligne)] = dosage
    # Prix
    ws["F" + str(nouvelle_ligne)] = pr_vent
    # Quantité
    ws["G" + str(nouvelle_ligne)] = quant
    # Médecin
    ws["H" + str(nouvelle_ligne)] = medecin
    # N° Ordonnance
    ws["I" + str(nouvelle_ligne)] = ordonnance
    # Nom Client
    ws["J" + str(nouvelle_ligne)] = nom_client
    # Sauvegarde
    wb.save(FICHIER_EXCEL)
    # Ajouter directement dans Treeview
    tableau_historique.insert("","end",values=(nouvelle_ligne - 1,numero_serie,
            datetime.now().strftime("%d/%m/%Y"),nom,dosage,pr_vent,quant,
            stock,medecin,ordonnance,nom_client))

    RECH.delete(0,"end")
    NOM.delete(0,"end")
    DOSAGE.delete(0,"end")
    STOCK.delete(0,"end")
    PR_VENT.delete(0,"end")
    QUAN.delete(0,"end")
    NOM_CLIENT.delete(0,"end")
    MEDECIN.delete(0,"end")
    N_ORD.delete(0,"end")

